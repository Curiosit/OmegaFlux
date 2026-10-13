"""
prepare_convergence_data.py

Cleans the Radiance convergence-check results (WO-A) into tidy, documented files
for the OmegaFlux repository.

Inputs (Grasshopper/Honeybee exports, paths configurable below or via CLI):
  new runs  : -ab 8 for rectangular, rotated and L-shaped rooms; -ab 12 for ten L-shapes
  originals : the -ab 5 exports used to build the training dataset
  selection : convergence_check/selected_rooms.csv (from select_convergence_rooms.py)

Outputs (OUT_DIR):
  convergence_sensors.csv.gz  one row per sensor; UDI fractions per setting (ab5 / ab8 / ab12)
  convergence_rooms.csv       one row per room; metadata, settings present, room means, shifts
  convergence_settings.json   simulation settings for each run
  convergence_cleaning_log.txt every row dropped or repaired, and why
  README_convergence.md       column dictionary

Usage (from the repository root):
  python -m omegaflux.prepare_convergence_data
  # new runs from data/raw/convergence/ (extracted from data/raw/convergence_runs.zip if needed),
  # original -ab 5 exports from data/ (as extracted by omegaflux.unzip_data)
"""
import argparse
import csv
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

csv.field_size_limit(sys.maxsize)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
FILES = {
    # (class, setting) -> new-run export
    ("rect", "ab8"):         "omegaflux-check-WO-A1c.xlsx",
    ("rect_rotated", "ab8"): "omegaflux-check-WO-A1rotated.xlsx",
    ("randomized", "ab8"):   "study-ml-12-Lshapes-corrected-second50study.xlsx",
    ("randomized", "ab12"):  "study-ml-12-Lshapes-corrected-second50study-ab12.xlsx",
}
ORIGINALS = {
    "rect":         "study-ml-rect.xlsx",
    "rect_rotated": "study-ml-rotated.xlsx",   # optional until provided
    "randomized":   "study-ml-shapes.csv",
}
SELECTION = os.path.join("data", "convergence_check", "selected_rooms.csv")
OUT_DIR = os.path.join("data", "convergence_check")
RUNS_DIR = os.path.join("data", "raw", "convergence")
RUNS_ZIP = os.path.join("data", "raw", "convergence_runs.zip")
ORIGINALS_DIR = "data"

# Rooms run at -ab 12 (ladder). Chosen after the -ab 8 results: the eight most-shifted
# L-shapes plus the two darkest. Any other -ab 12 row is excluded (e.g. off-by-one runs).
AB12_ROOMS = {f"Expert_LShape_{n}" for n in (98, 107, 127, 136, 224, 408, 439, 467, 850, 1043)}

SETTINGS = {
    "common": {
        "tool": "Ladybug Tools 1.10, HB Annual Daylight",
        "method": "enhanced 2-phase (direct sun traced to hourly solar positions; 145-patch Tregenza sky for diffuse)",
        "weather": "Warsaw Okecie EPW (same file as training data)",
        "udi_thresholds_lux": [100, 3000],
        "schedule": "component default, 8:00-18:00 daily",
        "sensor_grid": "0.25 m, 0.8 m above floor, identical to training data",
        "materials": "component defaults, identical to training data",
    },
    "ab5":  {"radiance_par": "-ab 5 -ad 5000 -lw 2e-05", "role": "original training labels"},
    "ab8":  {"radiance_par": "-ab 8 -ad 20000 -lw 5e-06", "role": "convergence reference"},
    "ab12": {"radiance_par": "-ab 12 -ad 20000 -lw 1e-06", "role": "ladder check, 10 L-shaped rooms"},
}

NAME_RE = re.compile(r"^(Expert_LShape_\d+|[A-D]_Rotated_\d+|[A-D]_\d+)$")
CLOSURE_TOL = 0.1   # pp, |low + useful + up - 100|
COORD_TOL = 1e-4    # m, sensor position agreement between runs
DECIMALS = 2

log_lines = []


def log(msg):
    log_lines.append(msg)
    print(msg)


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------
def nums(s, sep=";"):
    if s is None:
        return None
    vals = [v for v in str(s).split(sep) if v.strip() not in ("", "None")]
    return np.array(vals, dtype=float) if vals else None


def coords_from_matrix(s):
    """'{x, y, z}{x, y, z}...' -> (n, 2) array of x, y."""
    pts = re.findall(r"\{([^}]*)\}", str(s))
    return np.array([[float(v) for v in p.split(",")[:2]] for p in pts]) if pts else None


def coords_from_split(xs, ys):
    """'x1Xx2X...', 'y1Yy2Y...' -> (n, 2) array."""
    x = nums(xs, "X")
    y = nums(ys, "Y")
    if x is None or y is None or len(x) != len(y):
        return None
    return np.c_[x, y]


def room_name(rec):
    """The L-shape export has its name and index columns swapped; find the real name."""
    for key in ("Iteration_Name", "Iteration_Nr"):
        v = rec.get(key)
        if v is not None and NAME_RE.match(str(v).strip()):
            return str(v).strip()
    return None


def parse_record(rec):
    name = room_name(rec)
    if "Sensor_location_matrix X" in rec:
        xy = coords_from_split(rec.get("Sensor_location_matrix X"), rec.get("Sensor_location_matrix Y"))
    else:
        xy = coords_from_matrix(rec.get("Sensor_location_matrix"))
    return {
        "name": name,
        "timestamp": str(rec.get("DateTime")) if rec.get("DateTime") is not None else None,
        "xy": xy,
        "low": nums(rec.get("raw_UDI_low")),
        "useful": nums(rec.get("raw_UDI")),
        "up": nums(rec.get("raw_UDI_up")),
    }


def read_xlsx(path, wanted=None):
    ws = load_workbook(path, read_only=True).active
    it = ws.iter_rows(values_only=True)
    header = next(it)
    for row in it:
        if not any(c is not None for c in row):
            continue
        rec = dict(zip(header, row))
        if wanted is not None and room_name(rec) not in wanted:
            continue
        yield rec


def read_csv_rows(path, wanted=None):
    with open(path, newline="") as f:
        for rec in csv.DictReader(f):
            if wanted is not None and room_name(rec) not in wanted:
                continue
            yield rec


def read_any(path, wanted=None):
    return read_csv_rows(path, wanted) if path.lower().endswith(".csv") else read_xlsx(path, wanted)


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def validate(r, label):
    """Return a reason string if the record is unusable, else None."""
    if r["name"] is None:
        return "no recognisable room name"
    if any(r[k] is None for k in ("low", "useful", "up")):
        return "missing UDI results"
    n = len(r["useful"])
    if not (len(r["low"]) == len(r["up"]) == n):
        return f"UDI arrays differ in length ({len(r['low'])}/{n}/{len(r['up'])})"
    if r["xy"] is None or len(r["xy"]) != n:
        return f"sensor coordinates missing or length {0 if r['xy'] is None else len(r['xy'])} != {n}"
    dev = np.abs(r["low"] + r["useful"] + r["up"] - 100).max()
    if dev > CLOSURE_TOL:
        return f"fractions do not sum to 100 (max deviation {dev:.2f})"
    allv = np.concatenate([r["low"], r["useful"], r["up"]])
    if allv.min() < -1e-6 or allv.max() > 100 + 1e-6:
        return "values outside 0-100"
    return None


def load_runs(path, klass, setting, allowed, prefilter=False):
    """Parse one export; drop invalid rows, rooms outside `allowed`, and duplicates.
    prefilter=True skips unselected rooms silently (used for the full original exports)."""
    out = {}
    for rec in read_any(path, wanted=allowed if prefilter else None):
        r = parse_record(rec)
        tag = f"[{klass}/{setting}] {os.path.basename(path)}: {r['name'] or rec.get('Iteration_Nr')}"
        if r["name"] not in allowed:
            log(f"DROP {tag}: not in the selection for this run")
            continue
        reason = validate(r, tag)
        if reason:
            log(f"DROP {tag}: {reason}")
            continue
        if r["name"] in out:
            keep = max(out[r["name"]], r, key=lambda z: z["timestamp"] or "")
            log(f"DUPLICATE {tag}: kept run from {keep['timestamp']}")
            r = keep
        out[r["name"]] = r
    return out


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=RUNS_DIR, help="folder with the new-run exports")
    ap.add_argument("--originals", default=ORIGINALS_DIR, help="folder with the original -ab 5 exports")
    ap.add_argument("--selection", default=SELECTION)
    ap.add_argument("--out", default=OUT_DIR)
    args = ap.parse_args()
    P = lambda f: os.path.join(args.root, f)
    O = lambda f: os.path.join(args.originals, f)
    if not os.path.isdir(args.root) and os.path.exists(RUNS_ZIP):
        import zipfile
        with zipfile.ZipFile(RUNS_ZIP) as z:
            z.extractall(os.path.dirname(RUNS_ZIP))
        log(f"Extracted {RUNS_ZIP}")

    sel = pd.read_csv(args.selection)
    selected = {k: set(g["Iteration_Name"]) for k, g in sel.groupby("Source_Tag")}

    # ---- new runs ---------------------------------------------------------
    runs = {}  # (class, setting) -> {name: record}
    for (klass, setting), fname in FILES.items():
        path = P(fname)
        if not os.path.exists(path):
            log(f"MISSING input {fname}: {klass}/{setting} skipped")
            continue
        allowed = AB12_ROOMS if setting == "ab12" else selected.get(klass, set())
        runs[(klass, setting)] = load_runs(path, klass, setting, allowed)
        log(f"OK {klass}/{setting}: {len(runs[(klass, setting)])} rooms from {fname}")

    # ---- originals (-ab 5) for the same rooms -----------------------------
    for klass, fname in ORIGINALS.items():
        names = set().union(*[set(v) for (k, s), v in runs.items() if k == klass]) if runs else set()
        path = O(fname)
        if not names:
            continue
        if not os.path.exists(path):
            log(f"MISSING original {fname}: {klass} rooms will have no ab5 columns")
            continue
        runs[(klass, "ab5")] = load_runs(path, klass, "ab5", names, prefilter=True)
        log(f"OK {klass}/ab5: {len(runs[(klass, 'ab5')])} rooms from {fname}")

    # ---- assemble per-sensor table ----------------------------------------
    sensor_frames, room_rows = [], []
    classes = sorted({k for k, _ in runs})
    for klass in classes:
        settings_here = [s for s in ("ab5", "ab8", "ab12") if (klass, s) in runs]
        names = sorted(set().union(*[set(runs[(klass, s)]) for s in settings_here]),
                       key=lambda n: (n.split("_")[0], int(n.split("_")[-1])))
        for name in names:
            present = [s for s in settings_here if name in runs[(klass, s)]]
            ref = runs[(klass, present[0])][name]
            xy = ref["xy"]
            ok = True
            for s in present[1:]:
                other = runs[(klass, s)][name]["xy"]
                if len(other) != len(xy) or np.abs(other - xy).max() > COORD_TOL:
                    log(f"DROP [{klass}] {name}: sensor grid differs between {present[0]} and {s}")
                    ok = False
            if not ok:
                continue

            df = pd.DataFrame({
                "room_id": name,
                "class": klass,
                "sensor_idx": np.arange(len(xy)),
                "x_m": np.round(xy[:, 0], 6),
                "y_m": np.round(xy[:, 1], 6),
            })
            row = {"room_id": name, "class": klass, "group_id": f"{klass}:{name}",
                   "n_sensors": len(xy), "settings": "+".join(present)}
            for s in ("ab5", "ab8", "ab12"):
                for frac in ("low", "useful", "up"):
                    col = f"udi_{frac}_{s}"
                    if s in present:
                        df[col] = np.round(runs[(klass, s)][name][frac], DECIMALS)
                        row[f"mean_{col}"] = round(float(runs[(klass, s)][name][frac].mean()), 3)
                    else:
                        df[col] = np.nan
                if s in present and s != "ab5":
                    row[f"timestamp_{s}"] = runs[(klass, s)][name]["timestamp"]
            for a, b in (("ab5", "ab8"), ("ab8", "ab12")):
                if a in present and b in present:
                    d = runs[(klass, b)][name]["useful"] - runs[(klass, a)][name]["useful"]
                    row[f"useful_bias_{a}_to_{b}"] = round(float(d.mean()), 3)
                    row[f"useful_mae_{a}_to_{b}"] = round(float(np.abs(d).mean()), 3)
            sensor_frames.append(df)
            room_rows.append(row)

    sensors = pd.concat(sensor_frames, ignore_index=True)
    rooms = pd.DataFrame(room_rows)
    meta = sel.rename(columns={"Iteration_Name": "room_id", "Source_Tag": "class"})[
        ["room_id", "class", "stratum", "reason", "v2_flag"]].rename(
        columns={"reason": "selection_reason", "v2_flag": "preselected_v2"})
    rooms = rooms.merge(meta, on=["room_id", "class"], how="left")
    rooms["ab12_run"] = rooms["settings"].str.contains("ab12")

    # ---- write ------------------------------------------------------------
    os.makedirs(args.out, exist_ok=True)
    sensors.to_csv(os.path.join(args.out, "convergence_sensors.csv.gz"), index=False, compression="gzip")
    rooms.to_csv(os.path.join(args.out, "convergence_rooms.csv"), index=False)
    with open(os.path.join(args.out, "convergence_settings.json"), "w") as f:
        json.dump(SETTINGS, f, indent=2)

    summary = rooms.groupby(["class", "settings"]).size().rename("rooms").reset_index()
    log("\nSummary\n" + summary.to_string(index=False))
    log(f"\n{len(rooms)} rooms, {len(sensors):,} sensor rows written to {args.out}/")
    with open(os.path.join(args.out, "convergence_cleaning_log.txt"), "w") as f:
        f.write(f"Generated {dt.datetime.now().isoformat(timespec='seconds')}\n\n" + "\n".join(log_lines) + "\n")

    write_readme(args.out, rooms, sensors)


def write_readme(out, rooms, sensors):
    txt = f"""# Radiance convergence check

Re-simulation of {len(rooms)} held-out rooms with stricter Radiance settings, to test whether
the training labels (`-ab 5 -ad 5000 -lw 2e-05`) under-resolve interreflection.
Rooms were selected by `select_convergence_rooms.py` (stratified, held-out partition only);
the ten `-ab 12` rooms are the eight most-shifted L-shapes at `-ab 8` plus the two darkest.
Only the Radiance parameters differ between runs; geometry, sensors, weather, materials,
thresholds and schedule are identical (see `convergence_settings.json`).

## convergence_sensors.csv.gz ({len(sensors):,} rows, one per sensor)
| column | meaning |
|---|---|
| room_id | room name as in the training data (e.g. `B_373`, `C_Rotated_16`, `Expert_LShape_408`) |
| class | `rect`, `rect_rotated` or `randomized` (L-shaped) |
| sensor_idx | 0-based sensor order, identical across runs and to the training data |
| x_m, y_m | sensor position (m); sensors at 0.8 m above floor |
| udi_low_<s>, udi_useful_<s>, udi_up_<s> | % of occupied hours below 100 lux / 100-3000 lux / above 3000 lux, for setting `<s>` |

Settings: `ab5` original labels, `ab8` convergence reference, `ab12` ladder (L-shapes only). Empty = not run.

## convergence_rooms.csv (one per room)
Room metadata (`stratum`, `selection_reason`, `preselected_v2`), settings present, room-mean UDI per
setting, and per-room shifts in useful UDI: `useful_bias_a_to_b` (mean signed change, pp) and
`useful_mae_a_to_b` (mean absolute change, pp).

## Notes
- `convergence_cleaning_log.txt` lists every row excluded and why (missing results, rooms run by
  mistake outside the selection, duplicates).
- Values rounded to {DECIMALS} decimals; low + useful + up = 100 within {CLOSURE_TOL} pp at every sensor.
"""
    with open(os.path.join(out, "README_convergence.md"), "w") as f:
        f.write(txt)


if __name__ == "__main__":
    main()
